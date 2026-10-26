import io
import math
import uuid
import zipfile
from openpyxl import load_workbook, Workbook

FIELDS = ['name', 'uuid', 'os', 'cpu', 'ram_gb', 'disk_gb', 'ip', 'host', 'cluster', 'power', 'kind', 'group_name', 'application']
ALIASES = {
 'name': ['VM', 'Name'], 'uuid': ['VM UUID', 'UUID', 'Instance UUID'],
 'os': ['OS according to the configuration file', 'OS according to the VMware Tools', 'Guest OS', 'OS'],
 'cpu': ['CPUs', 'Num CPU', 'CPU'], 'ram_gb': ['Memory', 'Memory MB'],
 'disk_gb': ['Provisioned MiB', 'Total disk capacity MiB'], 'ip': ['Primary IP Address', 'IP Address'],
 'host': ['Host'], 'cluster': ['Cluster'], 'power': ['Powerstate', 'Power State']}

class ValidationError(ValueError): pass

def number(value, label, integer=False):
    if value in ('', None): return None
    try: n = float(value)
    except (TypeError, ValueError): raise ValidationError(f'{label}: debe ser numérico')
    if not math.isfinite(n) or n < 0 or n > 1e9 or (integer and not n.is_integer()):
        raise ValidationError(f'{label}: valor fuera de rango')
    return int(n) if integer else round(n, 3)

def validate(data):
    result = {k: str(data.get(k) or '').strip() for k in FIELDS}
    if not result['name'] or len(result['name']) > 200: raise ValidationError('Nombre requerido (máximo 200 caracteres)')
    if any(len(v) > 500 for v in result.values()): raise ValidationError('Un campo excede los 500 caracteres')
    for k in ('cpu', 'ram_gb', 'disk_gb'): result[k] = number(data.get(k), k, k == 'cpu')
    result['uuid'] = result['uuid'].lower()
    result['kind'] = result['kind'] or 'virtual'
    if result['kind'] not in ('virtual', 'physical'): raise ValidationError('Tipo de servidor inválido')
    return result

def workbook(raw):
    if len(raw) > 20 * 1024 * 1024: raise ValidationError('Archivo demasiado grande (máximo 20 MB)')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValidationError('El Excel descomprimido excede 100 MB')
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except ValidationError: raise
    except Exception: raise ValidationError('Se requiere un archivo .xlsx válido, sin contraseña')

def parse_rvtools(raw):
    book = workbook(raw)
    try:
        sheet_name = next((s for s in book.sheetnames if s.lower() == 'vinfo'), None)
        if not sheet_name: raise ValidationError('No se encontró la hoja vInfo')
        sheet = book[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(v or '').strip().lower() for v in next(rows, ())]
        mapping = {k: next((headers.index(a.lower()) for a in aliases if a.lower() in headers), None) for k, aliases in ALIASES.items()}
        if mapping['name'] is None: raise ValidationError('vInfo debe contener la columna VM')
        result, errors, seen = [], [], set()
        for line, row in enumerate(rows, 2):
            if line > 10001: raise ValidationError('Máximo 10,000 filas por archivo')
            if not any(v is not None for v in row): continue
            data = {k: row[i] if i is not None and i < len(row) else '' for k, i in mapping.items()}
            try:
                for k in ('ram_gb', 'disk_gb'):
                    n = number(data[k], k)
                    data[k] = round(n / 1024, 3) if n is not None else None
                data['kind'] = 'virtual'
                data = validate(data)
                if not data['uuid']: raise ValidationError('Falta VM UUID; se requiere para conciliar sin duplicados')
                if data['uuid'] in seen: raise ValidationError('UUID repetido en el archivo')
                seen.add(data['uuid'])
                result.append({'line': line, 'data': data})
            except ValidationError as e: errors.append({'line': line, 'error': str(e)})
        return result, errors
    finally: book.close()

def anonymize(source, target):
    """Allowlisted data only; never copy original workbook metadata or other sheets."""
    from pathlib import Path
    if Path(source).resolve() == Path(target).resolve(): raise ValidationError('El destino debe ser distinto al original')
    if Path(target).exists(): raise ValidationError('El destino ya existe; elige otro nombre')
    records, errors = parse_rvtools(Path(source).read_bytes())
    if errors: raise ValidationError(f'No se generó la muestra: {len(errors)} filas requieren corrección; primera: {errors[0]}')
    book = Workbook()
    sheet = book.active
    sheet.title = 'vInfo'
    sheet.append(['VM', 'VM UUID', 'OS', 'CPUs', 'Memory', 'Provisioned MiB', 'Primary IP Address', 'Host', 'Cluster', 'Powerstate'])
    host_map, cluster_map = {}, {}
    for i, item in enumerate(records[:100], 1):
        d = item['data']
        def token(mapping, value, prefix):
            if not value: return ''
            if value not in mapping: mapping[value] = f'{prefix}-{len(mapping)+1:03}'
            return mapping[value]
        # Fixed synthetic values prevent names hidden in free text, capacity fingerprints and formulas.
        sheet.append([f'srv-{i:03}', str(uuid.uuid4()), 'Linux de ejemplo', 2, 4096, 40960, f'192.0.2.{i}', token(host_map, d['host'], 'host'), token(cluster_map, d['cluster'], 'cluster'), 'poweredOn'])
    book.save(target)
    return min(len(records), 100)
