import io
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch
from urllib.request import Request, build_opener, HTTPCookieProcessor
from urllib.error import HTTPError
from http.cookiejar import CookieJar
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import server
from inventory import parse_rvtools, anonymize, ValidationError
from openpyxl import Workbook, load_workbook

def fixture(rows=None):
    book=Workbook(); sheet=book.active; sheet.title='vInfo'
    sheet.append(['VM','VM UUID','CPUs','Memory','Provisioned MiB','Primary IP Address','Host','Cluster','OS','Powerstate'])
    for row in rows or [['secret-srv','ABC-123',4,8192,102400,'10.1.1.1','private-host','private-cluster','RHEL 8','poweredOn']]: sheet.append(row)
    extra=book.create_sheet('Secrets');extra.append(['do-not-copy'])
    book.properties.creator='Secret company'
    stream=io.BytesIO();book.save(stream);return stream.getvalue()

class InventoryTests(unittest.TestCase):
    def test_units_and_invalid_uuid(self):
        records,errors=parse_rvtools(fixture())
        self.assertEqual(errors,[])
        self.assertEqual(records[0]['data']['ram_gb'],8)
        self.assertEqual(records[0]['data']['disk_gb'],100)
        self.assertEqual(records[0]['data']['uuid'],'abc-123')
        _, errors=parse_rvtools(fixture([['srv','',1,1024,1024]]))
        self.assertEqual(len(errors),1)
    def test_duplicate_and_negative(self):
        _,errors=parse_rvtools(fixture([['one','id',1,1024,1024],['two','id',1,1024,1024],['bad','id3',-1,0,0]]))
        self.assertEqual(len(errors),2)
    def test_anonymizer_allowlist(self):
        with tempfile.TemporaryDirectory() as folder:
            src=Path(folder)/'original.xlsx'; dst=Path(folder)/'sample.xlsx';raw=fixture();src.write_bytes(raw)
            self.assertEqual(anonymize(src,dst),1)
            self.assertEqual(src.read_bytes(),raw)
            book=load_workbook(dst)
            self.assertEqual(book.sheetnames,['vInfo'])
            values=str(list(book['vInfo'].values))+str(book.properties)
            for secret in ['secret-srv','10.1.1.1','private-host','private-cluster','Secret company','do-not-copy','ABC-123']:self.assertNotIn(secret,values)
            book.close()
            with self.assertRaises(ValidationError):anonymize(src,src)
            with self.assertRaises(ValidationError):anonymize(src,dst)
    def test_http_inventory_import_auth(self):
        with tempfile.TemporaryDirectory() as folder:
            server.DATA=Path(folder);server.KEY_FILE=Path(folder)/'secrets'/'key';server.init()
            salt='00'*16
            with server.db() as c:c.execute("INSERT INTO users(username,salt,password,role) VALUES(?,?,?,'admin')",('admin',salt,server.password_hash('test-password-123',salt)))
            http=server.ThreadingHTTPServer(('127.0.0.1',0),server.Handler)
            thread=threading.Thread(target=http.serve_forever,daemon=True);thread.start()
            base=f'http://127.0.0.1:{http.server_port}/api/'
            client=build_opener(HTTPCookieProcessor(CookieJar()))
            def request(path,data=None,csrf='',raw=False):
                req=Request(base+path,data=None if data is None else data if raw else json.dumps(data).encode(),headers={'Content-Type':'application/octet-stream' if raw else 'application/json','X-CSRF-Token':csrf})
                with client.open(req) as response:return json.load(response)
            try:
                with self.assertRaises(HTTPError) as err:request('servers')
                self.assertEqual(err.exception.code,401)
                session=request('login',{'username':'admin','password':'test-password-123'});csrf=session['csrf']
                with self.assertRaises(HTTPError) as err:request('servers',{'name':'bad'})
                self.assertEqual(err.exception.code,403)
                with patch('server.discover',return_value={'data':{'name':'found'},'warnings':[]}) as discover:
                    with self.assertRaises(HTTPError):request('discover',{'ip':'192.0.2.1'})
                    discover.assert_not_called()
                    found=request('discover',{'ip':'192.0.2.1','password':'do-not-store'},csrf)
                    self.assertEqual(found['data']['name'],'found')
                    self.assertEqual(request('servers'),[])
                preview=request('import/preview',fixture(),csrf,True)
                self.assertEqual(request('servers'),[])
                result=request('import/commit',{'token':preview['token']},csrf)
                self.assertEqual(result,{'created':1,'updated':0})
                with self.assertRaises(HTTPError):request('import/commit',{'token':preview['token']},csrf)
                row=request('servers')[0];row['group_name']='ERP';row['application']='Billing'
                row['ssh']={'username':'reader','password':'private-ssh-secret','port':22,'enabled':True}
                request('servers',row,csrf)
                public=request('servers')[0]
                self.assertTrue(public['credentials_configured'])
                self.assertNotIn('private-ssh-secret',json.dumps(public))
                self.assertNotIn('secret',public)
                with server.db() as c:
                    token=c.execute('SELECT secret FROM server_credentials').fetchone()[0]
                    self.assertNotIn('private-ssh-secret',token)
                def verify_saved(payload,*args):
                    self.assertEqual(payload['password'],'private-ssh-secret')
                    return {'data':{'name':'found'},'warnings':[]}
                with patch('server.discover',side_effect=verify_saved) as discover:
                    request('discover',{'server_id':row['id'],'ip':row['ip'],'port':22,'username':'reader','password':''},csrf)
                    self.assertEqual(discover.call_count,1)
                    with self.assertRaises(HTTPError):request('discover',{'server_id':row['id'],'ip':'192.0.2.99','port':22,'username':'reader','password':''},csrf)
                    self.assertEqual(discover.call_count,1)
                self.assertEqual(request('dashboard')['servers'][0]['state'],'pending')
                self.assertTrue(request('metrics/refresh',{'id':row['id']},csrf)['queued'])
                preview=request('import/preview',fixture(),csrf,True)
                result=request('import/commit',{'token':preview['token']},csrf)
                self.assertEqual(result,{'created':0,'updated':1})
                self.assertEqual(request('servers')[0]['group_name'],'ERP')
                self.assertEqual(len(request('servers')),1)
                preview=request('import/preview',fixture([['bad','',1,1,1]]),csrf,True)
                with self.assertRaises(HTTPError):request('import/commit',{'token':preview['token']},csrf)
                request('settings',{'interval_minutes':10,'retention_days':90},csrf)
                server.init()
                self.assertEqual(request('settings')['retention_days'],90)
                request('logout',{},csrf)
                with self.assertRaises(HTTPError):request('servers')
            finally:http.shutdown();http.server_close();thread.join()

if __name__=='__main__':unittest.main()
